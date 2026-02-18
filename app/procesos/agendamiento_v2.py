from fastapi import APIRouter, UploadFile, File
from fastapi.responses import StreamingResponse
from io import BytesIO
import pandas as pd
import unicodedata
from openpyxl.chart import BarChart, Reference, PieChart
from openpyxl.chart.label import DataLabelList

router = APIRouter(prefix="/agendamiento-v2", tags=["Agendamiento v2"])

# ==================================================
# UTILIDADES
# ==================================================
def normalizar_texto(texto):
    if pd.isna(texto):
        return ""
    texto = str(texto)
    texto = unicodedata.normalize("NFD", texto)
    texto = texto.encode("ascii", "ignore").decode("utf-8")
    texto = texto.upper().strip()
    texto = " ".join(texto.split())
    return texto


def normalizar_fecha(valor):
    if pd.isna(valor):
        return None

    valor = str(valor)

    # Manager → 2026-01-12 16:05:00
    if "-" in valor:
        return pd.to_datetime(valor, errors="coerce").date()

    # Bitrix → 12/01/2026
    return pd.to_datetime(valor, errors="coerce", dayfirst=True).date()


# ==================================================
# ENDPOINT
# ==================================================
@router.post("/procesar")
async def procesar_agendamiento_v2(
    manager_file: UploadFile = File(...),
    bitrix_file: UploadFile = File(...),
):
    # -----------------------
    # LEER ARCHIVOS
    # -----------------------
    df_manager = pd.read_excel(manager_file.file, header=1)
    df_bitrix = pd.read_excel(bitrix_file.file)

    df_manager.columns = df_manager.columns.str.strip().str.upper()
    df_bitrix.columns = df_bitrix.columns.str.strip().str.upper()

    # -----------------------
    # LIMPIAR MANAGER
    # -----------------------
    df_manager = df_manager[
        [
            "GUIA#",
            "PZ",
            "PAIS",
            "FECHA",
            "REMITENTE",
            "DESTINATARIO",
            "COMENTARIOS",
            "PESO",
            "TOTAL",
            "METODO PAGO",
        ]
    ].copy()

    df_manager["NOMBRE_NORM"] = df_manager["REMITENTE"].apply(normalizar_texto)
    df_manager["FECHA_NORM"] = df_manager["FECHA"].apply(normalizar_fecha)

    # -----------------------
    # LIMPIAR BITRIX
    # -----------------------
    df_bitrix = df_bitrix[
        [
            "CLIENTE",
            "FECHA DE AGENDA",
            "FECHA RECOGIDA",
            "ASESOR",
            "CIUDAD",
            "TIPO DE CLIENTE",
            "TIPO ENVIO",
            "VENTA CAJA",
            "CANTIDAD CAJAS",
        ]
    ].copy()

    df_bitrix["NOMBRE_NORM"] = df_bitrix["CLIENTE"].apply(normalizar_texto)
    df_bitrix["FECHA_NORM"] = df_bitrix["FECHA RECOGIDA"].apply(normalizar_fecha)

    # ==================================================
    # DETECTAR MULTIPLES EN BITRIX
    # ==================================================
    conteo_bitrix = (
        df_bitrix.groupby(["NOMBRE_NORM", "FECHA_NORM"])
        .size()
        .reset_index(name="COUNT")
    )

    df_bitrix = df_bitrix.merge(
        conteo_bitrix,
        on=["NOMBRE_NORM", "FECHA_NORM"],
        how="left",
    )

    # ==================================================
    # NIVEL 1 — EXACTO
    # ==================================================
    cruzados_exacto = df_manager.merge(
        df_bitrix[df_bitrix["COUNT"] == 1],
        on=["NOMBRE_NORM", "FECHA_NORM"],
        how="inner",
    )

    # ==================================================
    # NIVEL 2 — MULTIPLE_FECHA
    # ==================================================
    cruzados_multiple = df_manager.merge(
        df_bitrix[df_bitrix["COUNT"] > 1],
        on=["NOMBRE_NORM", "FECHA_NORM"],
        how="inner",
    )

    # ==================================================
    # EXCLUIR YA CRUZADOS
    # ==================================================
    guias_cruzadas = pd.concat(
        [
            cruzados_exacto[["GUIA#"]],
            cruzados_multiple[["GUIA#"]],
        ]
    ).drop_duplicates()

    manager_restante = df_manager[
        ~df_manager["GUIA#"].isin(guias_cruzadas["GUIA#"])
    ]

    # ==================================================
    # NIVEL 3 — SOLO NOMBRE
    # ==================================================
    conteo_nombres = df_bitrix["NOMBRE_NORM"].value_counts()
    nombres_unicos = conteo_nombres[conteo_nombres == 1].index

    df_bitrix_unicos = df_bitrix[df_bitrix["NOMBRE_NORM"].isin(nombres_unicos)]

    cruzados_nombre = manager_restante.merge(
        df_bitrix_unicos,
        on="NOMBRE_NORM",
        how="inner",
    )

    # ==================================================
    # NIVEL 4 — NO CRUZADO
    # ==================================================
    guias_finales = pd.concat(
        [
            guias_cruzadas,
            cruzados_nombre[["GUIA#"]],
        ]
    ).drop_duplicates()

    no_cruzados = df_manager[
        ~df_manager["GUIA#"].isin(guias_finales["GUIA#"])
    ].copy()

    # ==================================================
    # UNIR TODO
    # ==================================================
    df_final = pd.concat(
        [
            cruzados_exacto,
            cruzados_multiple,
            cruzados_nombre,
            no_cruzados,
        ],
        ignore_index=True,
    )

    # ==================================================
    # LIMPIEZA FINAL (archivo limpio para usuario)
    # ==================================================
    columnas_a_eliminar = [
        "CLIENTE",
        "NOMBRE_NORM",
        "FECHA_NORM",
        "COUNT",
        "FECHA_NORM_x",
        "FECHA_NORM_y",
    ]

    df_final = df_final.drop(columns=columnas_a_eliminar, errors="ignore")

    # ==================================================
    # DASHBOARD (TABLAS)
    # ==================================================
    df_por_asesor = (
        df_final.groupby("ASESOR", dropna=False)
        .agg(AGENDAMIENTOS=("GUIA#", "count"), TOTAL_DINERO=("TOTAL", "sum"))
        .reset_index()
    )

    df_por_tipo_cliente = (
        df_final.groupby("TIPO DE CLIENTE", dropna=False)
        .agg(AGENDAMIENTOS=("GUIA#", "count"), TOTAL_DINERO=("TOTAL", "sum"))
        .reset_index()
    )

    df_por_metodo_pago = (
        df_final.groupby("METODO PAGO", dropna=False)
        .agg(TOTAL_DINERO=("TOTAL", "sum"))
        .reset_index()
    )

    df_por_ciudad = (
        df_final.groupby("CIUDAD", dropna=False)
        .agg(AGENDAMIENTOS=("GUIA#", "count"), TOTAL_DINERO=("TOTAL", "sum"))
        .reset_index()
    )

    # ==================================================
    # EXPORTAR EXCEL + GRÁFICAS
    # ==================================================
    output = BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:

        # Hoja principal
        df_final.to_excel(writer, index=False, sheet_name="Agendamiento")

        # Hoja Dashboard (tablas)
        fila = 0
        for df in [
            df_por_asesor,
            df_por_tipo_cliente,
            df_por_metodo_pago,
            df_por_ciudad,
        ]:
            df.to_excel(writer, sheet_name="Dashboard", startrow=fila, index=False)
            fila += len(df) + 3

        # -----------------------
        # GRÁFICAS
        # -----------------------
        workbook = writer.book
        dashboard = workbook["Dashboard"]

        # Gráfica 1 — Agendamientos por Asesor
        chart1 = BarChart()
        chart1.title = "Agendamientos por Asesor"
        chart1.y_axis.title = "Cantidad"
        chart1.x_axis.title = "Asesor"

        data = Reference(dashboard, min_col=2, min_row=1, max_row=len(df_por_asesor) + 1)
        cats = Reference(dashboard, min_col=1, min_row=2, max_row=len(df_por_asesor) + 1)

        chart1.add_data(data, titles_from_data=True)
        chart1.set_categories(cats)
        chart1.dataLabels = DataLabelList(showVal=True)
        dashboard.add_chart(chart1, "M2")

        # Gráfica 2 — Dinero por Asesor
        chart2 = BarChart()
        chart2.type = "bar"
        chart2.title = "Dinero por Asesor"

        data = Reference(dashboard, min_col=3, min_row=1, max_row=len(df_por_asesor) + 1)
        chart2.add_data(data, titles_from_data=True)
        chart2.set_categories(cats)
        chart2.dataLabels = DataLabelList(showVal=True)
        dashboard.add_chart(chart2, "M18")

        # Gráfica 3 — Tipo de Cliente
        start = len(df_por_asesor) + 4
        end = start + len(df_por_tipo_cliente)

        chart3 = PieChart()
        chart3.title = "Tipo de Cliente"

        data = Reference(dashboard, min_col=3, min_row=start, max_row=end)
        cats = Reference(dashboard, min_col=1, min_row=start + 1, max_row=end)

        chart3.add_data(data, titles_from_data=True)
        chart3.set_categories(cats)
        chart3.dataLabels = DataLabelList(showPercent=True)
        dashboard.add_chart(chart3, "M34")

        # Gráfica 4 — Método de Pago
        start = end + 4
        end = start + len(df_por_metodo_pago)

        chart4 = PieChart()
        chart4.title = "Método de Pago"

        data = Reference(dashboard, min_col=2, min_row=start, max_row=end)
        cats = Reference(dashboard, min_col=1, min_row=start + 1, max_row=end)

        chart4.add_data(data, titles_from_data=True)
        chart4.set_categories(cats)
        chart4.dataLabels = DataLabelList(showPercent=True)
        dashboard.add_chart(chart4, "M52")

        # Gráfica 5 — Ciudad
        start = end + 4
        end = start + len(df_por_ciudad)

        chart5 = BarChart()
        chart5.title = "Agendamientos por Ciudad"

        data = Reference(dashboard, min_col=2, min_row=start, max_row=end)
        cats = Reference(dashboard, min_col=1, min_row=start + 1, max_row=end)

        chart5.add_data(data, titles_from_data=True)
        chart5.set_categories(cats)
        chart5.dataLabels = DataLabelList(showVal=True)
        dashboard.add_chart(chart5, "M70")

    output.seek(0)

    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=Agendamiento.xlsx"},
    )
